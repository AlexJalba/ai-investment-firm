"""Excel daily report generator using openpyxl."""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.portfolio.models import PortfolioSnapshot

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="D6E4F0")


def _header(ws, row: int, cols: list[str]) -> None:
    for c, val in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def write_excel_report(
    path: str,
    trade_date: str,
    snapshot: PortfolioSnapshot,
    start_of_day_value: float,
    filled_trades: list[dict],
    research_reports: list[dict],
    benchmark_pct: float | None,
    narrative: str,
) -> None:
    wb = openpyxl.Workbook()
    pnl = snapshot.total_value - start_of_day_value
    pnl_pct = pnl / start_of_day_value if start_of_day_value else 0

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20

    title_font = Font(size=14, bold=True, color="1F4E79")
    ws["A1"] = f"AI Investment Firm — Daily Report {trade_date}"
    ws["A1"].font = title_font
    ws.merge_cells("A1:B1")

    rows = [
        ("Start of Day Value", f"${start_of_day_value:,.2f}"),
        ("End of Day Value", f"${snapshot.total_value:,.2f}"),
        ("Daily P&L", f"${pnl:+,.2f} ({pnl_pct:+.2%})"),
        ("Benchmark (SPY)", f"{benchmark_pct:+.2%}" if benchmark_pct is not None else "N/A"),
        ("Cash", f"${snapshot.cash:,.2f}"),
        ("Trades Executed", str(len(filled_trades))),
    ]
    for i, (label, value) in enumerate(rows, 3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=value)
        if "P&L" in label:
            cell.font = Font(color="006100" if pnl >= 0 else "9C0006", bold=True)

    ws.cell(row=len(rows) + 4, column=1, value="Narrative").font = Font(bold=True)
    ws.cell(row=len(rows) + 5, column=1, value=narrative)
    ws.merge_cells(f"A{len(rows) + 5}:B{len(rows) + 5 + 8}")
    ws.cell(row=len(rows) + 5, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    # ── Holdings sheet ────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Holdings")
    cols = ["Ticker", "Shares", "Cost Basis", "Sector"]
    _header(ws2, 1, cols)
    for r, h in enumerate(snapshot.holdings, 2):
        ws2.cell(r, 1, h.ticker)
        ws2.cell(r, 2, h.shares)
        ws2.cell(r, 3, h.cost_basis)
        ws2.cell(r, 4, h.sector)
        if r % 2 == 0:
            for c in range(1, 5):
                ws2.cell(r, c).fill = ACCENT_FILL
    for c in range(1, 5):
        ws2.column_dimensions[get_column_letter(c)].width = 18

    # ── Trades sheet ──────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Trades")
    cols = ["Time", "Ticker", "Side", "Shares", "Fill Price", "Notional", "Commission", "Slippage"]
    _header(ws3, 1, cols)
    for r, t in enumerate(filled_trades, 2):
        ws3.cell(r, 1, str(t.get("executed_at", ""))[:19])
        ws3.cell(r, 2, t.get("ticker", ""))
        side_cell = ws3.cell(r, 3, t.get("side", ""))
        side_cell.font = Font(color="006100" if t.get("side") == "BUY" else "9C0006", bold=True)
        ws3.cell(r, 4, t.get("shares"))
        ws3.cell(r, 5, t.get("fill_price"))
        ws3.cell(r, 6, t.get("notional"))
        ws3.cell(r, 7, t.get("commission"))
        ws3.cell(r, 8, t.get("slippage"))
    for c in range(1, 9):
        ws3.column_dimensions[get_column_letter(c)].width = 16

    # ── Research sheet ────────────────────────────────────────────────────────
    ws4 = wb.create_sheet("Research")
    cols = ["Ticker", "Sentiment", "Recommendation", "Confidence", "Summary"]
    _header(ws4, 1, cols)
    for r, rep in enumerate(research_reports, 2):
        ws4.cell(r, 1, rep.get("ticker", ""))
        ws4.cell(r, 2, rep.get("sentiment", ""))
        ws4.cell(r, 3, rep.get("recommendation", ""))
        ws4.cell(r, 4, rep.get("confidence"))
        ws4.cell(r, 5, rep.get("summary", ""))
        ws4.cell(r, 5).alignment = Alignment(wrap_text=True)
    ws4.column_dimensions["E"].width = 60
    for c in range(1, 5):
        ws4.column_dimensions[get_column_letter(c)].width = 18

    wb.save(path)
